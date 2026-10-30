"""In-memory CSV, XLSX and audit JSON exports for finalized documents."""

from __future__ import annotations

import csv
import json
from dataclasses import asdict
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from formflow.domain import FieldKey, FieldName, FieldValue
from formflow.template import PRODUCTS
from formflow.validation import ValidatedDocument

_HEADERS = ("product_id", "product", "unit", "quantity", "unit_price", "line_total")


def _require_validated(document: ValidatedDocument) -> None:
    if not isinstance(document, ValidatedDocument):
        raise TypeError("Exports require a ValidatedDocument.")


def _item_rows(document: ValidatedDocument) -> list[dict[str, Any]]:
    fields = {field.key: field for field in document.fields}
    rows: list[dict[str, Any]] = []
    for product in PRODUCTS:
        quantity = fields[FieldKey(product.product_id, FieldName.QUANTITY)].value
        unit_price = fields[FieldKey(product.product_id, FieldName.UNIT_PRICE)].value
        line_total = (
            Decimal(quantity) * unit_price
            if isinstance(quantity, int) and isinstance(unit_price, Decimal)
            else None
        )
        rows.append(
            {
                "product_id": product.product_id,
                "product": product.name,
                "unit": product.unit,
                "quantity": quantity,
                "unit_price": unit_price,
                "line_total": line_total,
            }
        )
    return rows


def _plain(value: FieldValue | Decimal | None) -> str:
    if value is None:
        return ""
    return format(value, "f") if isinstance(value, Decimal) else str(value)


def to_csv(document: ValidatedDocument) -> bytes:
    _require_validated(document)
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=_HEADERS, lineterminator="\n")
    writer.writeheader()
    for row in _item_rows(document):
        writer.writerow({key: _plain(value) for key, value in row.items()})
    return buffer.getvalue().encode("utf-8")


def to_xlsx(document: ValidatedDocument) -> bytes:
    _require_validated(document)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Itens"
    sheet.append(["ID", "Produto", "Unidade", "Quantidade", "Preço unitário", "Total"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F4C5C")

    for row in _item_rows(document):
        sheet.append(
            [
                row["product_id"],
                row["product"],
                row["unit"],
                row["quantity"],
                float(row["unit_price"]) if row["unit_price"] is not None else None,
                float(row["line_total"]) if row["line_total"] is not None else None,
            ]
        )
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, 5).number_format = 'R$ #,##0.00'
        sheet.cell(row_index, 6).number_format = 'R$ #,##0.00'
    for index, width in enumerate((12, 28, 14, 14, 18, 18), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _json_value(value: FieldValue) -> int | str | None:
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def to_audit_json(document: ValidatedDocument) -> bytes:
    _require_validated(document)
    payload = {
        "document_id": document.document_id,
        "image_sha256": document.image_sha256,
        "finalized_at": document.finalized_at.isoformat(),
        "timings_ms": dict(document.timings_ms),
        "fields": [
            {
                "product_id": field.key.product_id,
                "field_name": field.key.field_name.value,
                "value": _json_value(field.value),
                "raw_text": field.raw_text,
                "confidence": field.confidence,
                "initial_source": field.initial_source.value,
                "final_source": field.final_source.value,
                "human_interacted": field.human_interacted,
            }
            for field in document.fields
        ],
        "events": [
            {
                **asdict(event),
                "key": {
                    "product_id": event.key.product_id,
                    "field_name": event.key.field_name.value,
                },
                "previous_value": _json_value(event.previous_value),
                "new_value": _json_value(event.new_value),
                "at": event.at.isoformat(),
                "initial_source": event.initial_source.value,
                "final_source": event.final_source.value,
            }
            for event in document.events
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
